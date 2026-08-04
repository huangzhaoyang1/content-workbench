"""数据分析服务：解析公众号后台导出的 CSV/Excel，产出概览、趋势、榜单与选题建议。

设计要点：
- 只依赖标准库 csv + openpyxl，不引入 pandas，保持环境轻量、启动快。
- 列名自动映射：中文/英文常见表头都能识别，识别不到就返回 mapping 让前端提示。
- 数据集存内存 + 落地 backend/data/analytics_last.json，后端重启后仍能看到上次分析。
"""
from __future__ import annotations

import csv
import datetime
import io
import json
import re
import threading
import uuid
from pathlib import Path

from ..config import DATA_DIR

_STORE_PATH = DATA_DIR / "analytics_last.json"
_LOCK = threading.Lock()
_DATASETS: dict[str, dict] = {}

# 列名候选（小写、去空格后匹配）。顺序即优先级。
_FIELD_ALIASES: dict[str, list[str]] = {
    "title": ["标题", "文章标题", "题目", "内容标题", "title", "文章", "图文标题"],
    "date": ["日期", "时间", "发布时间", "发表时间", "推送时间", "date", "publish_time", "发布日期"],
    "reads": ["阅读量", "阅读数", "总阅读人数", "总阅读次数", "阅读", "送达阅读率分母", "reads", "read", "views", "浏览量", "阅读人数"],
    "likes": ["在看", "在看数", "在看人数", "点赞", "点赞数", "赞", "likes", "like", "wow"],
    "shares": ["分享", "分享数", "转发", "转发数", "分享人数", "shares", "share", "forward"],
}

_MAX_ROWS = 20000  # 防御：超大文件截断，避免内存爆掉


class AnalyticsError(Exception):
    """数据解析/分析过程中的可预期错误，路由层转成 400。"""


# --------------------------------------------------------------------------
# 解析
# --------------------------------------------------------------------------
def _norm_header(s: str) -> str:
    return re.sub(r"[\s_\-()（）%]+", "", str(s or "")).strip().lower()


def _detect_mapping(headers: list[str]) -> dict[str, str | None]:
    """把原始表头映射到 title/date/reads/likes/shares 五个标准字段。"""
    normed = {_norm_header(h): h for h in headers}
    mapping: dict[str, str | None] = {}
    for field, aliases in _FIELD_ALIASES.items():
        hit = None
        # 1) 完全相等
        for a in aliases:
            na = _norm_header(a)
            if na in normed:
                hit = normed[na]
                break
        # 2) 包含匹配（如「总阅读量(次)」包含「阅读量」）
        if hit is None:
            for a in aliases:
                na = _norm_header(a)
                for nh, raw in normed.items():
                    if na and na in nh:
                        hit = raw
                        break
                if hit:
                    break
        mapping[field] = hit
    return mapping


def _to_number(v) -> float | None:
    """把 '1,234'、'12.3%'、'1.2万' 之类的值转成数字；失败返回 None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(",", "").replace("，", "")
    if not s:
        return None
    mult = 1.0
    if s.endswith("万"):
        mult, s = 10000.0, s[:-1]
    elif s.endswith("w") or s.endswith("W"):
        mult, s = 10000.0, s[:-1]
    s = s.rstrip("%")
    try:
        return float(s) * mult
    except ValueError:
        return None


_DATE_FORMATS = (
    "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日",
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M",
    "%m/%d/%Y", "%m-%d", "%Y%m%d",
)


def _to_date(v) -> str | None:
    """归一化成 YYYY-MM-DD；解析不了返回 None。"""
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s:
        return None
    for f in _DATE_FORMATS:
        try:
            dt = datetime.datetime.strptime(s, f)
            if f == "%m-%d":
                dt = dt.replace(year=datetime.date.today().year)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            continue
    # 兜底：抓 2024-01-02 / 2024/1/2 这种子串
    m = re.search(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", s)
    if m:
        try:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).strftime("%Y-%m-%d")
        except ValueError:
            return None
    return None


def _decode_csv(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030", "big5"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_csv(content: bytes) -> tuple[list[str], list[list]]:
    text = _decode_csv(content)
    sample = text[:4096]
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except Exception:
        pass
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for i, row in enumerate(reader):
        if i > _MAX_ROWS:
            break
        rows.append(row)
    if not rows:
        raise AnalyticsError("文件内容为空")
    headers = [str(h).strip() for h in rows[0]]
    return headers, rows[1:]


def _parse_excel(content: bytes) -> tuple[list[str], list[list]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise AnalyticsError("服务端缺少 openpyxl，无法解析 Excel，请改用 CSV") from e
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise AnalyticsError(f"Excel 解析失败：{type(e).__name__}") from e
    ws = wb[wb.sheetnames[0]]
    rows: list[list] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > _MAX_ROWS:
            break
        rows.append(list(row))
    wb.close()
    if not rows:
        raise AnalyticsError("表格内容为空")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return headers, rows[1:]


def parse_file(filename: str, content: bytes) -> dict:
    """解析上传文件，返回数据集字典（含 dataset_id）。"""
    name = (filename or "").lower()
    if not content:
        raise AnalyticsError("文件为空")
    if name.endswith((".xlsx", ".xlsm")):
        headers, raw_rows = _parse_excel(content)
    elif name.endswith(".xls"):
        raise AnalyticsError("暂不支持旧版 .xls，请另存为 .xlsx 或 CSV 后再上传")
    elif name.endswith((".csv", ".txt")):
        headers, raw_rows = _parse_csv(content)
    else:
        # 没有后缀就先按 CSV 试
        headers, raw_rows = _parse_csv(content)
    return _build_dataset(filename or "上传数据", headers, raw_rows)


def _build_dataset(filename: str, headers: list[str], raw_rows: list[list]) -> dict:
    mapping = _detect_mapping(headers)
    idx = {f: (headers.index(c) if c in headers else None) for f, c in mapping.items() if c}
    records: list[dict] = []
    for row in raw_rows:
        if not row or all((c is None or str(c).strip() == "") for c in row):
            continue

        def cell(field):
            i = idx.get(field)
            if i is None or i >= len(row):
                return None
            return row[i]

        rec = {
            "title": (str(cell("title")).strip() if cell("title") is not None else ""),
            "date": _to_date(cell("date")),
            "reads": _to_number(cell("reads")),
            "likes": _to_number(cell("likes")),
            "shares": _to_number(cell("shares")),
        }
        if not rec["title"] and rec["reads"] is None:
            continue  # 整行没有可用信息，跳过
        records.append(rec)

    if not records:
        raise AnalyticsError(
            "没有解析出有效数据行。请确认表格至少包含「标题」和「阅读量」两列（列名支持中英文常见写法）。"
        )

    dataset_id = uuid.uuid4().hex[:12]
    ds = {
        "dataset_id": dataset_id,
        "filename": filename,
        "headers": headers,
        "mapping": mapping,
        "records": records,
        "row_count": len(records),
        "has_date": any(r["date"] for r in records),
        "uploaded_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    with _LOCK:
        _DATASETS[dataset_id] = ds
        _persist(ds)
    return ds


# --------------------------------------------------------------------------
# 持久化
# --------------------------------------------------------------------------
def _persist(ds: dict) -> None:
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        Path(_STORE_PATH).write_text(
            json.dumps(ds, ensure_ascii=False), encoding="utf-8"
        )
    except Exception:
        pass  # 持久化失败不影响主流程


def _restore() -> dict | None:
    try:
        if not _STORE_PATH.exists():
            return None
        ds = json.loads(_STORE_PATH.read_text(encoding="utf-8"))
        if isinstance(ds, dict) and ds.get("records"):
            _DATASETS[ds["dataset_id"]] = ds
            return ds
    except Exception:
        pass
    return None


def get_dataset(dataset_id: str | None = None) -> dict | None:
    """取数据集：给了 id 就精确取，没给就取最近一次（含磁盘恢复）。"""
    if dataset_id:
        ds = _DATASETS.get(dataset_id)
        if ds:
            return ds
    if _DATASETS:
        return list(_DATASETS.values())[-1]
    return _restore()


def clear() -> None:
    with _LOCK:
        _DATASETS.clear()
        try:
            if _STORE_PATH.exists():
                _STORE_PATH.unlink()
        except Exception:
            pass


# --------------------------------------------------------------------------
# 示例数据
# --------------------------------------------------------------------------
_SAMPLE_TITLES = [
    "我用 AI 写了 30 天公众号，说几句大实话",
    "小白也能上手：一小时搭好自己的选题助手",
    "别再手动整理素材了，这套流程省我 2 小时",
    "AI 副业第一个月，我踩了这 5 个坑",
    "从 0 到 1：把 AI 变成你的内容流水线",
    "为什么你的 AI 文章没人看？问题出在选题",
    "实测 3 个免费 AI 工具，只有 1 个值得留",
    "做内容三个月，我把这件事想明白了",
    "一条视频涨粉 800，我复盘了这 4 个动作",
    "AI 提效不是玄学：给你一份可照抄清单",
    "我把公众号后台数据拉出来分析了一遍",
    "普通人做 AI 副业，最现实的三条路",
    "写了 50 篇后，我总结的开头公式",
    "这个笨办法，让我的完读率翻了一倍",
    "别学工具了，先把这个思路搞清楚",
    "关于 AI 变现，泼一盆冷水",
    "内容选题库怎么建？我的做法很土但有效",
    "从阅读量 200 到 2000，中间发生了什么",
    "AI 生成的内容，怎么改才像人写的",
    "我的日更流程全公开（含踩坑记录）",
]


def build_sample() -> dict:
    """生成一份可用于演示的示例数据（20 篇，近 40 天）。"""
    import random

    rnd = random.Random(20260804)
    today = datetime.date.today()
    headers = ["日期", "标题", "阅读量", "在看", "分享"]
    rows: list[list] = []
    for i, title in enumerate(_SAMPLE_TITLES):
        d = today - datetime.timedelta(days=(len(_SAMPLE_TITLES) - i) * 2)
        base = rnd.randint(280, 2600)
        # 让「踩坑 / 实测 / 复盘」类标题表现更好，方便建议模块出有意义的结论
        if any(k in title for k in ("踩坑", "实测", "复盘", "大实话")):
            base = int(base * 1.7)
        reads = base
        likes = max(1, int(reads * rnd.uniform(0.012, 0.055)))
        shares = max(1, int(reads * rnd.uniform(0.008, 0.042)))
        rows.append([d.strftime("%Y-%m-%d"), title, reads, likes, shares])
    ds = _build_dataset("示例数据.csv", headers, rows)
    ds["is_sample"] = True
    return ds


# --------------------------------------------------------------------------
# 分析
# --------------------------------------------------------------------------
def _safe_avg(values: list[float]) -> float | None:
    vals = [v for v in values if v is not None]
    return (sum(vals) / len(vals)) if vals else None


def _filter_by_range(records: list[dict], time_range: str) -> list[dict]:
    """近7天/近30天/全部。没有日期的记录在按范围筛选时会被排除。"""
    if time_range in ("全部", "", None):
        return records
    days = {"近7天": 7, "近30天": 30, "近90天": 90}.get(time_range)
    if not days:
        return records
    dated = [r for r in records if r.get("date")]
    if not dated:
        return records
    # 以数据集中最新日期为基准，避免历史数据全被筛没
    latest = max(r["date"] for r in dated)
    try:
        latest_dt = datetime.datetime.strptime(latest, "%Y-%m-%d").date()
    except ValueError:
        return records
    start = latest_dt - datetime.timedelta(days=days - 1)
    out = []
    for r in dated:
        try:
            d = datetime.datetime.strptime(r["date"], "%Y-%m-%d").date()
        except ValueError:
            continue
        if start <= d <= latest_dt:
            out.append(r)
    return out


def analyze(ds: dict, time_range: str = "全部") -> dict:
    """产出概览、趋势、三类榜单。"""
    records = _filter_by_range(ds["records"], time_range)
    reads = [r["reads"] for r in records if r["reads"] is not None]
    likes = [r["likes"] for r in records if r["likes"] is not None]
    shares = [r["shares"] for r in records if r["shares"] is not None]

    overview = {
        "total_articles": len(records),
        "total_reads": int(sum(reads)) if reads else 0,
        "avg_reads": round(_safe_avg(reads), 1) if reads else None,
        "avg_likes": round(_safe_avg(likes), 1) if likes else None,
        "avg_shares": round(_safe_avg(shares), 1) if shares else None,
        "max_reads": int(max(reads)) if reads else None,
    }

    # 趋势：按日期聚合
    trend: list[dict] = []
    if ds.get("has_date"):
        buckets: dict[str, dict] = {}
        for r in records:
            d = r.get("date")
            if not d:
                continue
            b = buckets.setdefault(d, {"date": d, "reads": 0.0, "likes": 0.0, "shares": 0.0, "count": 0})
            b["reads"] += r["reads"] or 0
            b["likes"] += r["likes"] or 0
            b["shares"] += r["shares"] or 0
            b["count"] += 1
        trend = [
            {
                "date": k,
                "reads": int(v["reads"]),
                "likes": int(v["likes"]),
                "shares": int(v["shares"]),
                "count": v["count"],
            }
            for k, v in sorted(buckets.items())
        ]

    # 榜单
    def top(key: str, n: int = 10) -> list[dict]:
        items = [r for r in records if r.get(key) is not None]
        items.sort(key=lambda r: r[key], reverse=True)
        return [_rank_item(r, i) for i, r in enumerate(items[:n])]

    def top_like_rate(n: int = 10) -> list[dict]:
        items = []
        for r in records:
            if r.get("reads") and r["reads"] > 0 and r.get("likes") is not None:
                rr = dict(r)
                rr["like_rate"] = round(r["likes"] / r["reads"] * 100, 2)
                items.append(rr)
        items.sort(key=lambda r: r["like_rate"], reverse=True)
        return [_rank_item(r, i) for i, r in enumerate(items[:n])]

    rankings = {
        "reads": top("reads"),
        "shares": top("shares"),
        "like_rate": top_like_rate(),
    }

    return {
        "dataset_id": ds["dataset_id"],
        "filename": ds.get("filename"),
        "time_range": time_range,
        "has_date": bool(ds.get("has_date")),
        "row_count": ds.get("row_count"),
        "mapping": ds.get("mapping"),
        "overview": overview,
        "trend": trend,
        "rankings": rankings,
    }


def _rank_item(r: dict, i: int) -> dict:
    reads = int(r["reads"]) if r.get("reads") is not None else None
    likes = int(r["likes"]) if r.get("likes") is not None else None
    # 在看率：三个榜单都补齐，前端表格才能按它排序、横向对比
    rate = r.get("like_rate")
    if rate is None and reads and reads > 0 and likes is not None:
        rate = round(likes / reads * 100, 2)
    return {
        "rank": i + 1,
        "title": r.get("title") or "（无标题）",
        "date": r.get("date"),
        "reads": reads,
        "likes": likes,
        "shares": int(r["shares"]) if r.get("shares") is not None else None,
        "like_rate": rate,
    }


# --------------------------------------------------------------------------
# 选题建议
# --------------------------------------------------------------------------
# 关键词 → (方向名, 建议切入角度)
_KEYWORD_THEMES: list[tuple[tuple[str, ...], str, str]] = [
    (("踩坑", "翻车", "坑", "失败", "教训"), "踩坑复盘型内容",
     "把一次真实翻车经历写成可复用的避坑清单，开头直接抛结果，中间讲过程，结尾给三条checklist。"),
    (("实测", "对比", "测评", "值得", "哪个好"), "工具实测与横向对比",
     "选 2-3 个同类工具做同一件事，给出耗时/效果/成本三维对比表，结论明确到「谁该用哪个」。"),
    (("复盘", "总结", "个月", "天", "从0", "从 0"), "阶段性复盘与数据公开",
     "用时间线讲清楚关键转折点，公开真实数据（哪怕不好看），结尾沉淀方法论。"),
    (("小白", "入门", "一小时", "上手", "教程", "怎么做"), "小白向保姆级教程",
     "抛开术语，用生活化类比讲清概念，配可照抄的步骤，每步给出预期结果截图位。"),
    (("变现", "副业", "赚钱", "收入", "接单"), "变现路径与真实收益",
     "讲清楚一条具体路径的完整链路：从哪来流量、怎么转化、实际赚多少、门槛在哪。"),
    (("效率", "提效", "流程", "自动化", "省"), "流程自动化与提效",
     "以「省下多少时间」为钩子，拆解自动化流程图，给出可直接复制的配置。"),
]


def suggestions(ds: dict, analysis: dict | None = None) -> list[dict]:
    """基于榜单数据生成 3-5 条方向建议，每条含依据与切入角度。"""
    a = analysis or analyze(ds, "全部")
    ov = a["overview"]
    top_reads = a["rankings"]["reads"]
    top_shares = a["rankings"]["shares"]
    top_rate = a["rankings"]["like_rate"]
    out: list[dict] = []

    # 1) 关键词主题命中：统计 TOP 榜单标题里的关键词
    pool = [t["title"] for t in top_reads[:8]] + [t["title"] for t in top_shares[:5]]
    for keys, name, angle in _KEYWORD_THEMES:
        hits = [t for t in pool if any(k in t for k in keys)]
        if len(hits) >= 2:
            sample = hits[0][:22]
            out.append({
                "name": name,
                "evidence": f"高表现文章里有 {len(hits)} 篇属于这一类，例如《{sample}…》，说明读者对这个方向反馈明显更好。",
                "angle": angle,
            })
        if len(out) >= 3:
            break

    # 2) 头部文章带来的方向建议
    if top_reads:
        best = top_reads[0]
        avg = ov.get("avg_reads") or 0
        ratio = (best["reads"] / avg) if (avg and best.get("reads")) else None
        ev = f"《{best['title'][:22]}…》以 {best['reads']} 阅读排第一"
        if ratio and ratio >= 1.3:
            ev += f"，是平均值的 {ratio:.1f} 倍，明显跑赢大盘"
        out.append({
            "name": f"复用爆款结构：{best['title'][:16]}",
            "evidence": ev + "。同题材换角度再写一轮，是性价比最高的选题来源。",
            "angle": "保留原文的选题内核，换一个身份视角或时间维度重写（比如从「我怎么做」换成「三个月后回头看」）。",
        })

    # 3) 传播效率：分享榜与阅读榜差异
    if top_shares and top_reads:
        share_only = [t["title"] for t in top_shares[:5] if t["title"] not in [x["title"] for x in top_reads[:5]]]
        if share_only:
            out.append({
                "name": "高转发选题：工具清单与资源整理",
                "evidence": f"《{share_only[0][:22]}…》分享数进前五但阅读没进前五，说明它的「收藏/转发价值」高于标题吸引力。",
                "angle": "做成可保存的清单体（工具表、话术库、避坑表），标题直接标注「收藏向」，提高二次传播。",
            })

    # 4) 在看率：内容共鸣度
    if top_rate:
        best_rate = top_rate[0]
        out.append({
            "name": "情绪共鸣型内容",
            "evidence": f"《{best_rate['title'][:22]}…》在看率 {best_rate.get('like_rate')}%，远高于普通篇目，读者愿意为它「表态」。",
            "angle": "在文章中段加入一段带态度的判断或反常识观点，结尾抛一个立场问题，引导读者点在看。",
        })

    # 5) 数据量太小时给个稳妥建议
    if ov["total_articles"] < 5:
        out.append({
            "name": "先把数据样本做厚",
            "evidence": f"当前只有 {ov['total_articles']} 篇数据，规律还不稳定，结论容易被单篇噪声带偏。",
            "angle": "先按固定节奏产出 10 篇以上，同一栏目连续做，再回来做横向对比分析。",
        })

    # 去重 + 截断到 5 条
    seen = set()
    uniq = []
    for s in out:
        if s["name"] in seen:
            continue
        seen.add(s["name"])
        uniq.append(s)
    return uniq[:5]
